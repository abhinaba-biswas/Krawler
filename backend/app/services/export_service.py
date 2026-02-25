"""
Export service – converts crawl results to CSV / XLSX / Apple Numbers.

Column layout mirrors the Gulf Food demo file:
  name | hall | stand | website | facebook | instagram | youtube
"""
from __future__ import annotations

import csv
import io
import re
from typing import Any, Dict, List, Optional
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.crawl import CrawlResult
from app.utils.logging import get_logger

log = get_logger(__name__)

COLUMNS = ["name", "hall", "stand", "website", "facebook", "instagram", "youtube"]

# ── Social-link patterns ────────────────────────────────────────────────────────

_FB = re.compile(
    r"https?://(?:www\.)?facebook\.com/[^\s\"'<>]+", re.IGNORECASE
)
_IG = re.compile(
    r"https?://(?:www\.)?instagram\.com/[^\s\"'<>]+", re.IGNORECASE
)
_YT = re.compile(
    r"https?://(?:www\.)?(?:youtube\.com|youtu\.be)/[^\s\"'<>]+", re.IGNORECASE
)


def _first(pattern: re.Pattern, sources: List[str]) -> Optional[str]:
    for src in sources:
        if not src:
            continue
        m = pattern.search(src)
        if m:
            return m.group(0).rstrip(".,;)")
    return None


def _extract_social(result: CrawlResult) -> Dict[str, Optional[str]]:
    """Pull facebook / instagram / youtube from links and raw HTML."""
    links: List[str] = result.links or []
    md = result.content_markdown or ""
    html = result.content_html or ""
    sd = result.structured_data or {}

    # Build a combined haystack: links list + markdown text + structured data dump
    haystack = links + [md, html, str(sd)]

    return {
        "facebook": _first(_FB, haystack),
        "instagram": _first(_IG, haystack),
        "youtube": _first(_YT, haystack),
    }


def _extract_hall_stand(result: CrawlResult) -> Dict[str, Optional[str]]:
    """Try to extract hall / stand from structured_data or metadata."""
    sd: Dict[str, Any] = result.structured_data or {}
    meta: Dict[str, Any] = result.page_metadata or {}

    hall = (
        sd.get("hall")
        or meta.get("hall")
        or sd.get("location")
        or meta.get("location")
    )
    stand = (
        sd.get("stand")
        or meta.get("stand")
        or sd.get("stand_number")
        or meta.get("stand_number")
    )
    return {"hall": hall, "stand": stand}


def _result_to_row(result: CrawlResult) -> Dict[str, Any]:
    meta: Dict[str, Any] = result.page_metadata or {}
    social = _extract_social(result)
    loc = _extract_hall_stand(result)

    return {
        "name": result.title or meta.get("og_title") or result.url,
        "hall": loc["hall"],
        "stand": loc["stand"],
        "website": result.url,
        "facebook": social["facebook"],
        "instagram": social["instagram"],
        "youtube": social["youtube"],
    }


# ── DB fetch ────────────────────────────────────────────────────────────────────

async def fetch_all_results(job_id: UUID, db: AsyncSession) -> List[CrawlResult]:
    stmt = (
        select(CrawlResult)
        .where(CrawlResult.job_id == job_id)
        .order_by(CrawlResult.timestamp)
    )
    return list((await db.execute(stmt)).scalars().all())


# ── Serialisers ─────────────────────────────────────────────────────────────────

def to_csv_bytes(rows: List[Dict[str, Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=COLUMNS)
    writer.writeheader()
    for row in rows:
        writer.writerow({c: row.get(c) or "" for c in COLUMNS})
    return buf.getvalue().encode("utf-8-sig")   # utf-8-sig adds BOM for Excel compat


def to_xlsx_bytes(rows: List[Dict[str, Any]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Crawl Results"

    # Header style
    header_font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    header_fill = PatternFill("solid", fgColor="1F3864")   # dark navy
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    # Data rows
    alt_fill = PatternFill("solid", fgColor="EBF0FA")
    for idx, row in enumerate(rows):
        ws.append([row.get(c) or "" for c in COLUMNS])
        row_num = idx + 2
        for cell in ws[row_num]:
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if idx % 2 == 1:
                cell.fill = alt_fill

    # Auto-width (max 60 chars)
    for col_idx, col_name in enumerate(COLUMNS, 1):
        max_len = max(
            len(col_name),
            *(len(str(row.get(col_name) or "")) for row in rows)
        ) if rows else len(col_name)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 60)

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_numbers_bytes(rows: List[Dict[str, Any]]) -> bytes:
    """
    Generate an Apple Numbers file using numbers-parser.
    """
    import numbers_parser

    doc = numbers_parser.Document()
    sheet = doc.sheets[0]
    sheet.name = "Crawl Results"
    table = sheet.tables[0]

    # Write headers
    for col_idx, col_name in enumerate(COLUMNS):
        table.write(0, col_idx, col_name)

    # Write data
    for row_idx, row in enumerate(rows, start=1):
        for col_idx, col_name in enumerate(COLUMNS):
            val = row.get(col_name) or ""
            table.write(row_idx, col_idx, val)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── Public API ─────────────────────────────────────────────────────────────────

async def export_job(
    job_id: UUID,
    fmt: str,
    db: AsyncSession,
) -> bytes:
    """Return the serialised export bytes for the given format (csv/xlsx/numbers)."""
    results = await fetch_all_results(job_id, db)
    rows = [_result_to_row(r) for r in results]

    fmt = fmt.lower()
    if fmt == "csv":
        return to_csv_bytes(rows)
    elif fmt == "xlsx":
        return to_xlsx_bytes(rows)
    elif fmt == "numbers":
        return to_numbers_bytes(rows)
    else:
        raise ValueError(f"Unsupported export format: {fmt!r}")
